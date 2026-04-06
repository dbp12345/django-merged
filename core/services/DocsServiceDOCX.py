from copy import deepcopy
from docx import Document
from pathlib import Path
from django.conf import settings
from io import BytesIO
from django.http import HttpResponse


class DocsServiceDOCX:
    # doc_path = Path(settings.BASE_DIR) / "core" / "documents_templates" / "manifests" / "Hand Crew Manifest 2024 T2C (Federal)_sample.docx"
    doc_path = Path(settings.BASE_DIR) / "core" / "documents_templates" / "manifests" / "Hand Crew Manifest 2.0.docx"
    output_path = Path(settings.BASE_DIR) / "core" / "documents_templates" / "manifests" / "filled_manifest.docx"

    def __init__(self):
        self.context = {}
        self.employees = []
        self.vehicles = []

    def setContext(self, context: dict):
        self.context = context

    def setEmployeesList(self, employees: list[dict]):
        self.employees = employees

    def setVehiclesList(self, vehicles: list[dict]):
        self.vehicles = vehicles

    def _render_doc(self):
        doc = Document(str(self.doc_path))

        self._replace_placeholders_in_paragraphs(doc)

        self.fill_employee_table(doc)
        self.fill_vehicle_table(doc)

        self._replace_placeholders_in_tables(doc, self.context)

        return doc

    def render_docx_to_file(self, output_path: str | Path = None):
        output_path = output_path or self.output_path
        doc = self._render_doc()
        doc.save(str(output_path))

    def generate_docx_response(self, filename = "filled_manifest") -> HttpResponse:
        doc = self._render_doc()
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.docx"'
        return response

    def _replace_placeholders_in_paragraphs(self, doc: Document):
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                run_text = run.text
                for key, val in self.context.items():
                    run_text = run_text.replace(f"{{{{{key}}}}}", str(val))
                run.text = self._remove_unmatched_aliases(run_text)

    def _replace_placeholders_in_tables(self, doc: Document, context: dict):
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text
                    for key, val in context.items():
                        cell_text = cell_text.replace(f"{{{{{key}}}}}", str(val))
                    cell.text = self._remove_unmatched_aliases(cell_text)

    def fill_employee_table(self, doc: Document):
        employee_table, template_row = self._find_template_row(doc, "{{emp_nun}}")
        if not employee_table:
            raise ValueError("Table employee not found")

        for emp in self.employees:
            new_row = deepcopy(template_row._tr)
            employee_table._tbl.append(new_row)
            row_cells = employee_table.rows[-1].cells
            for cell in row_cells:
                cell_text = cell.text
                for key, val in emp.items():
                    cell_text = cell_text.replace(f"{{{{{key}}}}}", str(val))
                cell.text = self._remove_unmatched_aliases(cell_text)

        employee_table._tbl.remove(template_row._tr)

    def fill_vehicle_table(self, doc: Document):
        vehicle_table, template_row = self._find_template_row(doc, "{{veh_drv_1}}")
        if not vehicle_table:
            raise ValueError("Table vehicle not found")

        for veh in self.vehicles:
            new_row = deepcopy(template_row._tr)
            vehicle_table._tbl.append(new_row)
            row_cells = vehicle_table.rows[-1].cells
            for cell in row_cells:
                cell_text = cell.text
                for key, val in veh.items():
                    cell_text = cell_text.replace(f"{{{{{key}}}}}", str(val))
                cell.text = self._remove_unmatched_aliases(cell_text)

        vehicle_table._tbl.remove(template_row._tr)

    def _find_template_row(self, doc: Document, marker: str):
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if marker in cell.text:
                        return table, row
        return None, None

    def _remove_unmatched_aliases(self, text: str) -> str:
        while "{{" in text and "}}" in text:
            start = text.find("{{")
            end = text.find("}}", start)
            if end == -1:
                break
            text = text[:start] + "" + text[end + 2:]
        return text
